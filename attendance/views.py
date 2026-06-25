from django.db import transaction
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import JsonResponse
from django.views.decorators.http import require_POST
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
import calendar
from datetime import datetime, date, timedelta
import json
import math  # [필수] 거리 계산용 라이브러리 (누락 방지)
from django.db.models import Q, Sum

# [필수] 공휴일 라이브러리
try:
    import holidays
except ImportError:
    holidays = None

# 모델 Import
from django.contrib.auth.models import User  # ★ 추가
from accounts.models import Profile, Process, Cohort, PartLeader
from quiz.models import StudentLog # [필수] 알림 로그용
from .models import WorkType, DailySchedule, ScheduleRequest, Attendance 


# ------------------------------------------------------------------
# [Helper] 연차 발생 개수 계산 함수 (근속연수 기준)
# ------------------------------------------------------------------
def calculate_annual_leave_total(profile, target_year):
    """
    입사일(joined_at) 기준으로 해당 연도의 총 연차 개수를 계산합니다.
    """
    if not profile.joined_at:
        return 15  # 입사일 없으면 기본값
    
    # 근속 연수 계산 (대상 년도 - 입사 년도)
    years_worked = target_year - profile.joined_at.year
    
    if years_worked < 1:
        return 15  # 1년차 미만
    
    # 가산 연차 계산: (근속연수 - 1) // 2
    added_days = (years_worked - 1) // 2
    if added_days < 0: added_days = 0
    
    total = 15 + int(added_days)
    
    # 최대 25개 제한 (근로기준법)
    return min(total, 25)


# ------------------------------------------------------------------
# [Helper] 스케줄 수정 권한 확인
# ------------------------------------------------------------------
def can_manage_schedule(user, target_profile):
    if user.is_superuser:
        return True
    
    if hasattr(user, 'profile') and user.profile.is_manager:
        if user.profile.process == target_profile.process:
            return True
            
    return False


# ------------------------------------------------------------------
# [Helper] 거리 계산 함수 (Haversine 공식)
# ------------------------------------------------------------------
def calculate_distance(lat1, lon1, lat2, lon2):
    """Haversine 공식을 이용한 거리 계산 (단위: km)"""
    R = 6371  # 지구 반지름 (km)
    d_lat = math.radians(lat2 - lat1)
    d_lon = math.radians(lon2 - lon1)
    a = math.sin(d_lat / 2) * math.sin(d_lat / 2) + \
        math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * \
        math.sin(d_lon / 2) * math.sin(d_lon / 2)
    c = 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
    return R * c


# ------------------------------------------------------------------
# 1. [신규] 스마트 출근 인증 (GPS + MDM/Camera Block)
# ------------------------------------------------------------------
@login_required
@require_POST
def process_attendance(request):
    """
    [신규] 출근 인증 처리 (AJAX 요청)
    - 프론트엔드에서 1차 검증(GPS, 카메라 차단) 후 넘어온 데이터 저장
    """
    try:
        # 1. 오늘 이미 출근했는지 확인
        today = timezone.localdate()
        if Attendance.objects.filter(user=request.user, date=today).exists():
             return JsonResponse({'status': 'fail', 'message': '이미 오늘의 출근 기록이 존재합니다.'})

        # 2. 출근 기록 저장
        Attendance.objects.create(
            user=request.user,
            date=today,
            check_in_time=timezone.now(),
            status='출근', 
            is_verified=True # 인증 성공 표시
        )
        
        return JsonResponse({'status': 'success', 'message': '출근 인증이 완료되었습니다! 오늘도 화이팅하세요.'})

    except Exception as e:
        return JsonResponse({'status': 'error', 'message': f'시스템 오류: {str(e)}'})

# (구버전 호환용)
@login_required
def upload_mdm(request):
    return redirect('attendance:mdm_status') 

@login_required
def mdm_status(request):
    return render(request, 'attendance/index.html') 


# ------------------------------------------------------------------
# 2. 캘린더 스케줄 조회 (매니저 전체조회 가능 / 교육생 보안 유지)
# ------------------------------------------------------------------
@login_required
def schedule_index(request):
    today = timezone.localdate()
    try:
        year = int(request.GET.get('year', today.year))
        month = int(request.GET.get('month', today.month))
    except ValueError:
        year, month = today.year, today.month
    
    kr_holidays = holidays.KR(years=year) if holidays else {}
    _, num_days = calendar.monthrange(year, month)
    days_in_month = []
    weekday_map = {0:'월', 1:'화', 2:'수', 3:'목', 4:'금', 5:'토', 6:'일'}

    for day in range(1, num_days + 1):
        d = date(year, month, day)
        days_in_month.append({
            'day': day, 
            'date_str': d.strftime('%Y-%m-%d'),
            'weekday': weekday_map[d.weekday()],
            'is_weekend': d.weekday() >= 5,
            'is_holiday': d in kr_holidays,
            'holiday_name': kr_holidays.get(d, ''),
            'is_today': d == today
        })

    user = request.user
    
    # 1. 쿼리셋 기본 준비 (재직 중인 인원)
    profiles = Profile.objects.select_related('cohort', 'process', 'user').filter(
        status='attending'
    ).exclude(name__isnull=True).exclude(name='')

    # 2. 관리자/매니저 권한 여부 확인
    is_manager_or_admin = user.is_superuser or (
        hasattr(user, 'profile') and (user.profile.is_manager or user.profile.is_pl)
    )

    # 3. [핵심 로직] 기수 및 공정 필터링
    # request.GET이 비어있으면 '처음 접속(Initial Load)' -> 자동 추천 동작
    # request.GET이 있으면(?cohort= 등) '검색 동작(Search)' -> 사용자가 선택한 값 존중 (빈값이면 전체조회)
    is_initial_load = (len(request.GET) == 0)

    sel_cohort = request.GET.get('cohort', '')
    sel_process = request.GET.get('process', '')
    sel_role = request.GET.get('role', 'student')

    # (A) 매니저/관리자 로직
    if is_manager_or_admin:
        if is_initial_load:
            # [처음 접속 시] 편의를 위해 '현재 활동 기수'와 '내 공정' 자동 선택
            # 1순위: 오늘 포함된 기수
            active_cohort = Cohort.objects.filter(
                start_date__lte=today, 
                end_date__gte=today
            ).first()

            if active_cohort:
                sel_cohort = str(active_cohort.id)
            else:
                # 2순위: 최신 기수
                latest_cohort = Cohort.objects.order_by('-start_date').first()
                if latest_cohort:
                    sel_cohort = str(latest_cohort.id)

            # 내 공정 자동 선택
            if hasattr(user, 'profile') and user.profile.process:
                sel_process = str(user.profile.process.id)
        
        # [검색 시] 위 자동 선택 로직을 타지 않으므로, sel_cohort가 비어있으면 '전체 조회'가 됨

        # 역할 필터
        if sel_role == 'manager':
            profiles = profiles.filter(
                Q(is_manager=True) | Q(is_pl=True) | 
                Q(user__is_superuser=True) | Q(user__is_staff=True)
            )
        else:
            profiles = profiles.filter(
                is_manager=False, is_pl=False, 
                user__is_superuser=False, user__is_staff=False
            )
        
        # 필터 적용 (값이 있을 때만 필터링 -> 값이 없으면 전체 조회)
        if sel_cohort:
            profiles = profiles.filter(cohort_id=sel_cohort)
        if sel_process:
            profiles = profiles.filter(process_id=sel_process)
        
    # (B) 교육생 로직 (보안 필수)
    else:
        # [중요] 교육생은 본인 기수/공정 강제 고정 (타 데이터 조회 불가)
        sel_role = 'student'
        
        if hasattr(user, 'profile'):
            my_profile = user.profile
            
            # 1. 내 기수만
            if my_profile.cohort:
                profiles = profiles.filter(cohort=my_profile.cohort)
                sel_cohort = str(my_profile.cohort.id) # 템플릿 표시용
            
            # 2. 내 공정만
            if my_profile.process:
                profiles = profiles.filter(process=my_profile.process)
                sel_process = str(my_profile.process.id) # 템플릿 표시용
            else:
                # 공정이 없으면 본인 데이터만
                profiles = profiles.filter(user=user)
                
            # 관리자 제외하고 순수 학생만
            profiles = profiles.filter(is_manager=False, is_pl=False, user__is_superuser=False)
        else:
            profiles = profiles.none()

    view_scope = request.GET.get('view', 'team')
    if view_scope == 'self' and hasattr(user, 'profile'):
        profiles = profiles.filter(user=user)

    profiles = profiles.order_by('name')

    # [연차 및 스케줄 데이터 매핑]
    current_year_start = date(year, 1, 1)
    current_year_end = date(year, 12, 31)

    leave_usage_map = {}
    if profiles.exists():
        usage_data = DailySchedule.objects.filter(
            profile__in=profiles,
            date__range=(current_year_start, current_year_end)
        ).values('profile').annotate(used_total=Sum('work_type__deduction'))

        for item in usage_data:
            leave_usage_map[item['profile']] = item['used_total'] or 0

    schedule_map = {}
    start_date = date(year, month, 1)
    end_date = date(year, month, num_days)
    
    db_data = {}
    if profiles.exists():
        schedules = DailySchedule.objects.filter(
            profile__in=profiles, date__range=(start_date, end_date)
        ).select_related('work_type')

        for s in schedules:
            if s.profile_id not in db_data: db_data[s.profile_id] = {}
            db_data[s.profile_id][s.date.strftime('%Y-%m-%d')] = s.work_type

    for p in profiles:
        total_leave = calculate_annual_leave_total(p, year)
        used = leave_usage_map.get(p.id, 0)
        remain = total_leave - used
        
        row_data = {
            'profile': p, 
            'daily_data': {}, 
            'stats': {
                'work':0, 'rest':0, 'leave':0, 'half':0, 'etc':0,
                'annual_remain': remain,
                'annual_total': total_leave
            }
        }
        user_schedules = db_data.get(p.id, {})
        
        for day_info in days_in_month:
            d_str = day_info['date_str']
            if d_str in user_schedules:
                wt = user_schedules[d_str]
                row_data['daily_data'][d_str] = wt
                if wt.deduction == 1.0: row_data['stats']['leave'] += 1
                elif 0 < wt.deduction < 1.0: row_data['stats']['half'] += 1
                elif wt.is_working_day and wt.deduction == 0: row_data['stats']['work'] += 1
                else:
                    if not wt.is_working_day: row_data['stats']['rest'] += 1
                    else: row_data['stats']['etc'] += 1
            else:
                if day_info['is_weekend'] or day_info['is_holiday']:
                    row_data['daily_data'][d_str] = None
                    row_data['stats']['rest'] += 1 
                else:
                    row_data['daily_data'][d_str] = 'DEFAULT_F' 
                    row_data['stats']['work'] += 1
                    
        schedule_map[p.id] = row_data

    # 다음달 계산
    if today.month == 12:
        next_month_start = date(today.year + 1, 1, 1)
    else:
        next_month_start = date(today.year, today.month + 1, 1)

    # 본인 profile_id 추출 (교육생 본인 행만 클릭 허용)
    my_profile_id = ''
    if hasattr(user, 'profile'):
        try:
            my_profile_id = user.profile.id
        except Exception:
            my_profile_id = ''

    context = {
        'year': year, 'month': month,
        'days_in_month': days_in_month,
        'schedule_map': schedule_map,
        'work_types': WorkType.objects.all().order_by('order'),
        'cohorts': Cohort.objects.all().order_by('-start_date'),
        'processes': Process.objects.all(),
        'sel_cohort': int(sel_cohort) if sel_cohort else '',
        'sel_process': int(sel_process) if sel_process else '',
        'sel_role': sel_role,
        'prev_month': (start_date - timedelta(days=1)).strftime('%Y-%m'),
        'next_month': (end_date + timedelta(days=1)).strftime('%Y-%m'),
        'is_manager': is_manager_or_admin,
        'my_profile_id': my_profile_id,  # ★ 추가
        'view_scope': view_scope,
    }
    return render(request, 'attendance/schedule.html', context)

# ------------------------------------------------------------------
# 3. 스케줄 수정 (기존 로직 유지)
# ------------------------------------------------------------------
@login_required
@require_POST
def update_schedule(request):
    try:
        data         = json.loads(request.body)
        profile_id   = data.get('profile_id')
        date_str     = data.get('date')
        work_type_id = data.get('work_type_id')
        reason       = (data.get('reason') or '').strip()

        target_profile = get_object_or_404(Profile, pk=profile_id)
        work_type      = get_object_or_404(WorkType, pk=work_type_id)
        target_date    = datetime.strptime(date_str, '%Y-%m-%d').date()
        today          = timezone.localdate()

        from quiz.models import Notification
        from django.urls import reverse

        # ── 역할 판별 ──────────────────────────────────────────────
        is_superuser = request.user.is_superuser

        # 본인 여부
        is_owner = False
        try:
            is_owner = (target_profile.user_id == request.user.pk)  # ★ id 직접 비교 (역참조 오류 방지)
        except Exception:
            pass

        # 대상이 매니저/PL/관리자인지 확인
        target_is_manager = (
            getattr(target_profile, 'is_manager', False) or
            getattr(target_profile, 'is_pl', False) or
            target_profile.user.is_superuser or
            target_profile.user.is_staff
        )

        # 매니저 권한:
        #   - 본인 행 제외 (is_owner=False)
        #   - 대상이 순수 교육생인 경우만 (target_is_manager=False)
        #   - 같은 공정
        is_manager_of_target = False
        if not is_superuser and not is_owner and hasattr(request.user, 'profile'):
            rp = request.user.profile
            if getattr(rp, 'is_manager', False) or getattr(rp, 'is_pl', False):
                if rp.process == target_profile.process and not target_is_manager:  # ★ 매니저→매니저 차단
                    is_manager_of_target = True

        # 기본 접근 차단
        if not (is_owner or is_superuser or is_manager_of_target):
            return JsonResponse({'status': 'error', 'message': '수정 권한이 없습니다.'})

        # ══════════════════════════════════════════════════════════
        # ① 최종관리자: 날짜/대상 무관 직접 저장
        # ══════════════════════════════════════════════════════════
        if is_superuser:
            DailySchedule.objects.update_or_create(
                profile=target_profile, date=target_date,
                defaults={'work_type': work_type}
            )
            return JsonResponse({'status': 'success', 'message': '관리자 권한으로 저장되었습니다.'})

        # ══════════════════════════════════════════════════════════
        # ② 과거 날짜: 최종관리자 외 전면 차단
        # ══════════════════════════════════════════════════════════
        if target_date < today:
            return JsonResponse({'status': 'error', 'message': '지난 날짜는 최종 관리자만 수정할 수 있습니다.'})

        # 다음달 1일
        if today.month == 12:
            next_month_start = date(today.year + 1, 1, 1)
        else:
            next_month_start = date(today.year, today.month + 1, 1)

        # 알림 발송 헬퍼
        def send_notifications(recipients, msg):
            for recipient in recipients:
                Notification.objects.create(
                    recipient=recipient,
                    sender=request.user,
                    message=msg,
                    # ★ 기존 주소 맨 뒤에 &show_pending=true 를 딱 붙여줍니다!
                    related_url=reverse('attendance:schedule_index') + f"?year={target_date.year}&month={target_date.month}&show_pending=true",
                    icon='bi-calendar-event',
                    notification_type='general'
                )

        superusers = User.objects.filter(is_superuser=True)

        def send_to_superusers(msg):
            send_notifications(superusers, msg)

        def send_to_managers_and_superusers(msg):
            managers = User.objects.filter(
                profile__is_manager=True,
                profile__process=target_profile.process
            )
            recipients = (managers | superusers).distinct()
            send_notifications(recipients, msg)

        def create_request_and_notify(msg_func):
            if not reason:
                return JsonResponse({'status': 'reason_required'})
            ScheduleRequest.objects.create(
                requester=target_profile, date=target_date,
                target_work_type=work_type, reason=reason, status='pending'
            )
            msg_func()
            return JsonResponse({'status': 'request_sent', 'message': '승인 요청이 전송되었습니다. 승인 후 반영됩니다.'})

        def direct_save():
            DailySchedule.objects.update_or_create(
                profile=target_profile, date=target_date,
                defaults={'work_type': work_type}
            )
            return JsonResponse({'status': 'success', 'message': '근무가 저장되었습니다.'})

        # ══════════════════════════════════════════════════════════
        # ③ 매니저 → 담당 교육생 행
        #    당일 → 최종관리자 결재 / 당월~익월 → 직접저장
        # ══════════════════════════════════════════════════════════
        if is_manager_of_target:
            if target_date == today:
                # 당일: 최종관리자 결재
                return create_request_and_notify(lambda: send_to_superusers(
                    f"📅 [당일 근무변경 요청] {request.user.get_full_name() or request.user.username} 매니저가 {target_profile.name}님의 {target_date.strftime('%m/%d')} '{work_type.short_name}' 변경을 요청했습니다."
                ))
            else:
                # 당월(당일제외) ~ 익월: 직접저장
                return direct_save()

        # ══════════════════════════════════════════════════════════
        # ④ 본인 행 수정 (매니저 본인 or 교육생 본인)
        # ══════════════════════════════════════════════════════════
        if is_owner:
            req_user_is_manager = hasattr(request.user, 'profile') and (
                getattr(request.user.profile, 'is_manager', False) or
                getattr(request.user.profile, 'is_pl', False)
            )

            if req_user_is_manager:
                # 매니저 본인
                if target_date >= next_month_start:
                    # 익월: 직접저장
                    return direct_save()
                else:
                    # 당일 + 당월(당일제외): 최종관리자 결재
                    return create_request_and_notify(lambda: send_to_superusers(
                        f"📅 [매니저 본인 근무변경 요청] {target_profile.name}님이 {target_date.strftime('%m/%d')} '{work_type.short_name}' 변경을 요청했습니다."
                    ))
            else:
                # 교육생 본인
                if target_date == today:
                    # 당일: 최종관리자만 결재
                    return create_request_and_notify(lambda: send_to_superusers(
                        f"📅 [교육생 당일 근무변경 요청] {target_profile.name}님이 {target_date.strftime('%m/%d')} '{work_type.short_name}' 변경을 요청했습니다. (당일 건 - 최종관리자 결재 필요)"
                    ))
                else:
                    # 당월(당일제외) + 익월: 매니저 + 최종관리자 결재
                    return create_request_and_notify(lambda: send_to_managers_and_superusers(
                        f"📅 [근무변경 요청] {target_profile.name}님이 {target_date.strftime('%m/%d')} '{work_type.short_name}' 변경을 요청했습니다."
                    ))

        return JsonResponse({'status': 'error', 'message': '처리할 수 없는 요청입니다.'})

    except Exception as e:
        import traceback
        print(f"❌ update_schedule 오류:\n{traceback.format_exc()}")
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

@login_required
def get_pending_requests(request):
    if request.user.is_superuser:
        # 최종관리자: 모든 대기 요청
        reqs = ScheduleRequest.objects.filter(status='pending').select_related(
            'requester', 'target_work_type'
        ).order_by('-id') # ★ 여기에 최신순 정렬 추가!
    elif hasattr(request.user, 'profile') and (
        request.user.profile.is_manager or request.user.profile.is_pl
    ):
        # 매니저: 담당 공정 교육생 요청만
        my_process = request.user.profile.process
        reqs = ScheduleRequest.objects.filter(
            requester__process=my_process,
            requester__is_manager=False,
            requester__is_pl=False,
            status='pending'
        ).select_related('requester', 'target_work_type').order_by('-id') # ★ 여기도 최신순 정렬 추가!
    else:
        return JsonResponse({'requests': []})
    data = [
        {
            'id': r.id,
            'name': r.requester.name,
            'date': r.date.strftime('%Y-%m-%d'),
            'type': r.target_work_type.short_name,
            'reason': r.reason,
            # ★ 역할 정보 추가
            'role': '교수' if (
                getattr(r.requester, 'is_manager', False) or
                getattr(r.requester, 'is_pl', False)
            ) else '교육생',
        }
        for r in reqs
    ]
    return JsonResponse({'requests': data})


@login_required
def process_request(request, request_id, action):
    """
    근무 변경 요청 승인/반려 통합 처리
    """
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': '잘못된 접근입니다.'}, status=405)

    if not request.user.is_staff:
        return JsonResponse({'status': 'error', 'message': '관리자 권한이 없습니다.'})

    try:
        with transaction.atomic():
            req_obj = get_object_or_404(ScheduleRequest, id=request_id)
            
            if req_obj.status != 'pending':
                return JsonResponse({'status': 'error', 'message': '이미 처리된 요청입니다.'})

            # 승인 (Approve)
            if action == 'approve':
                req_obj.status = 'approved'
                req_obj.approver = request.user
                req_obj.save()

                daily, created = DailySchedule.objects.get_or_create(
                    profile=req_obj.requester,  
                    date=req_obj.date
                )
                daily.work_type = req_obj.target_work_type
                daily.save()

                StudentLog.objects.create(
                    profile=req_obj.requester,
                    log_type='others',
                    reason=f"[근무변경 승인] {req_obj.date} 근무가 '{req_obj.target_work_type.short_name}'(으)로 변경되었습니다.",
                    is_resolved=True,
                    recorder=request.user
                )
                return JsonResponse({'status': 'success', 'message': '승인이 완료되었습니다.'})

            # 반려 (Reject)
            elif action == 'reject':
                req_obj.status = 'rejected'
                req_obj.approver = request.user
                req_obj.save()

                # ★ 1. 패널티가 적용되지 않도록 log_type을 'others'(기타)로 변경
                StudentLog.objects.create(
                    profile=req_obj.requester,
                    log_type='others',
                    reason=f"[근무변경 반려] {req_obj.date} 요청이 반려되었습니다. (사유: {req_obj.reason})",
                    is_resolved=True,
                    recorder=request.user
                )
                
                # ★ 2. 학생이 알 수 있도록 우측 상단 종소리 알림(Notification) 추가 발송
                from quiz.models import Notification
                Notification.objects.create(
                    recipient=req_obj.requester.user,
                    sender=request.user,
                    notification_type='general',
                    message=f"📅 {req_obj.date.strftime('%m/%d')} 근무 변경 요청이 반려되었습니다.",
                    related_url="/attendance/" # 학생 근태 페이지 연결 (필요시 URL 수정)
                )

                return JsonResponse({'status': 'success', 'message': '요청이 반려되었습니다.'})

            else:
                return JsonResponse({'status': 'error', 'message': '알 수 없는 명령입니다.'})

    except Exception as e:
        print(f"❌ [에러발생] process_request 중 오류: {e}")
        return JsonResponse({'status': 'error', 'message': f'서버 오류 발생: {str(e)}'}, status=500)


# ------------------------------------------------------------------
# 5. 전체 정상 적용 (일괄 처리)
# ------------------------------------------------------------------
@login_required
@require_POST
def apply_all_normal(request):
    """평일 일괄 적용"""
    try:
        data = json.loads(request.body)
        year = int(data.get('year'))
        month = int(data.get('month'))
        profile_ids = data.get('profile_ids', [])
        
        if not (request.user.is_superuser or (hasattr(request.user, 'profile') and request.user.profile.is_manager)):
             return JsonResponse({'status': 'error', 'message': '권한이 없습니다.'})

        normal_type = WorkType.objects.filter(name__contains="정상", deduction=0).first()
        if not normal_type: 
            normal_type = WorkType.objects.filter(deduction=0).exclude(name__contains="연차").order_by('order').first()
            
        if not normal_type:
            return JsonResponse({'status': 'error', 'message': '정상 근무 유형이 없습니다.'})
        
        kr_holidays = holidays.KR(years=year) if holidays else {}
        _, num_days = calendar.monthrange(year, month)
        create_list = []
        
        my_process = request.user.profile.process if hasattr(request.user, 'profile') else None
        
        for pid in profile_ids:
            target_profile = Profile.objects.get(pk=pid)
            if not request.user.is_superuser:
                if target_profile.process != my_process:
                    continue 

            for day in range(1, num_days + 1):
                curr_date = date(year, month, day)
                if curr_date.weekday() >= 5 or curr_date in kr_holidays:
                    continue

                if not DailySchedule.objects.filter(profile_id=pid, date=curr_date).exists():
                    create_list.append(DailySchedule(profile_id=pid, date=curr_date, work_type=normal_type))
        
        DailySchedule.objects.bulk_create(create_list)
        return JsonResponse({'status': 'success', 'count': len(create_list)})
        
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
    

@login_required
def check_in_page(request):
    """출근 체크 화면"""
    return render(request, 'attendance/check_in.html')

@login_required
@csrf_exempt
def check_in_api(request):
    """출근 체크 API (수정됨)"""
    if request.method != 'POST':
        return JsonResponse({'status': 'error', 'message': '잘못된 접근입니다.'})

    try:
        data = json.loads(request.body)
        lat = float(data.get('lat', 0))
        lon = float(data.get('lon', 0))
        is_mdm_active = data.get('is_mdm_active', False)

        # 1. MDM 검사
        if not is_mdm_active:
            return JsonResponse({'status': 'fail', 'message': '보안 앱(MDM)이 감지되지 않았습니다. 카메라가 차단되었는지 확인해주세요.'})

        # 2. 위치 검사 (300m)
        CENTER_LAT = 37.039159  # ⚠️ 실제 위도로 수정 필요
        CENTER_LON = 127.060482 # ⚠️ 실제 경도로 수정 필요
        RADIUS_LIMIT = 0.3

        distance = calculate_distance(lat, lon, CENTER_LAT, CENTER_LON)
        
        if distance > RADIUS_LIMIT:
            return JsonResponse({
                'status': 'fail', 
                'message': f'사업장 반경 {int(RADIUS_LIMIT*1000)}m 이내에서만 출근 가능합니다.\n(현재 거리: {int(distance*1000)}m)'
            })

        # 3. 출근 기록 (★ 여기가 수정되었습니다 ★)
        today = timezone.localdate()

        # daily_schedule 대신 user와 date로 중복 검사
        if Attendance.objects.filter(user=request.user, date=today).exists():
            return JsonResponse({'status': 'fail', 'message': '이미 금일 출근 기록이 있습니다.'})

        # daily_schedule 없이 직접 저장
        Attendance.objects.create(
            user=request.user,           # ★ user 저장
            date=today,                  # ★ date 저장
            check_in_time=timezone.now(),
            status='출근',               # 상태 (한글 or 영문 통일 필요)
            is_verified=True             # MDM/GPS 통과했으므로 인증됨
        )

        return JsonResponse({'status': 'success', 'message': '출근 인증이 완료되었습니다!'})

    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)})
    

@login_required
def export_schedule_excel(request):
    import openpyxl
    import calendar
    from django.http import HttpResponse
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from datetime import date

    # 1. 고급 모달 전용 파라미터 수집
    target_group = request.GET.get('target_group', 'all')  # manager, student, all
    download_scope = request.GET.get('scope', 'month')    # month, year
    
    # 🎯 [추가] 대시보드 연동형 정밀 필터 파라미터 수집
    req_cohort = request.GET.get('cohort', '')
    req_process = request.GET.get('process', '')
    req_company = request.GET.get('company', '')

    try:
        target_year = int(request.GET.get('year', timezone.now().year))
        target_month = int(request.GET.get('month', timezone.now().month))
    except ValueError:
        target_year, target_month = timezone.now().year, timezone.now().month

    # ── [★ 수정/추가]: 기수 선택 시 기수의 전체 운영 기간(예: 6월~7월)을 자동 추적하는 방어선 구축 ──
    if req_cohort:
        try:
            cohort_obj = Cohort.objects.get(id=req_cohort)
            start_date = cohort_obj.start_date
            end_date = cohort_obj.end_date
            target_year = start_date.year
        except:
            start_date = date(target_year, target_month, 1)
            _, num_days = calendar.monthrange(target_year, target_month)
            end_date = date(target_year, target_month, num_days)
    else:
        if download_scope == 'year':
            start_date = date(target_year, 1, 1)
            end_date = date(target_year, 12, 31)
        else:
            start_date = date(target_year, target_month, 1)
            _, num_days = calendar.monthrange(target_year, target_month)
            end_date = date(target_year, target_month, num_days)
    # ───────────────────────────────────────────────────────────────────────────────────

    # 2. 대상 프로필 셋 기본 필터링 (재직 중인 인원)
    profiles = Profile.objects.filter(status='attending').exclude(name__isnull=True).exclude(name='')
    
    # ══════════════════════════════════════════════════════════
    # 🔒 [보안 가드 소스] 최종관리자(Superuser)가 아니면 본인 공정 데이터만 강제 고정
    # ══════════════════════════════════════════════════════════
    if not request.user.is_superuser:
        if hasattr(request.user, 'profile') and request.user.profile.process:
            # 주소창 파라미터를 강제로 위조해도 백엔드에서 세션 유저의 공정 ID로 덮어써서 철저히 방어합니다.
            req_process = str(request.user.profile.process.id)
            profiles = profiles.filter(process_id=req_process)
        else:
            # 스태프 권한은 있으나 공정이 배정되지 않은 예외 케이스는 데이터 전면 차단
            profiles = profiles.none()
    else:
        # 최종 관리자는 선택한 공정이 있을 때만 필터링 (빈값이면 전체 조회 존중)
        if req_process:
            profiles = profiles.filter(process_id=req_process)
    # ══════════════════════════════════════════════════════════

    # 🎯 [추가] 기수 및 회사 동적 필터 처리 (값이 존재할 때만 스캔 작동)
    if req_cohort:
        profiles = profiles.filter(cohort_id=req_cohort)
    if req_company:
        profiles = profiles.filter(company_id=req_company) # 만약 company가 외래키가 아니라 텍스트면 company=req_company로 변경

    # 대상 그룹 필터 (교수용 / 교육생용 / 전체용)
    if target_group == 'manager':
        profiles = profiles.filter(Q(is_manager=True) | Q(is_pl=True) | Q(user__is_superuser=True) | Q(user__is_staff=True))
    elif target_group == 'student':
        profiles = profiles.filter(is_manager=False, is_pl=False, user__is_superuser=False, user__is_staff=False)
    
    profiles = profiles.order_by('name')

    # 3. 엑셀 워크북 바인딩 가동
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # 기본 시트 제거

    # ── [★ 수정/위치이동]: 아랫줄에 있던 스타일 사전 정의를 여기로 올립니다 ──
    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'), right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'), bottom=Side(style='thin', color='D3D3D3')
    )
    header_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    sat_fill = PatternFill(start_color='E6F2FF', end_color='E6F2FF', fill_type='solid') # 토요일 하늘색
    sun_fill = PatternFill(start_color='FFE6E6', end_color='FFE6E6', fill_type='solid') # 일요일 분홍색
    # ──────────────────────────────────────────────────────────────────

    # ──────────────────────────────────────────────────────────
    # ➕ [신규 추가]: Sheet1 첫 번째 칸에 연차/공가 사용 목록리스트 생성
    # ──────────────────────────────────────────────────────────
    ws_summary = wb.create_sheet(title="연차사용 목록리스트")
    ws_summary.column_dimensions['A'].width = 12
    ws_summary.column_dimensions['B'].width = 15
    ws_summary.column_dimensions['C'].width = 15
    ws_summary.column_dimensions['D'].width = 15
    ws_summary.column_dimensions['E'].width = 40

    # 요약 리스트 시트 헤더 디자인
    summary_headers = ["이름", "공정", "날짜", "근무 형태", "비고"]
    for ci, h in enumerate(summary_headers, 1):
        cell = ws_summary.cell(row=1, column=ci, value=h)
        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1F3864', end_color='1F3864', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border
    ws_summary.row_dimensions[1].height = 26

    # ── [★ 수정]: 전체 시작일과 종료일 연동 ──
    all_start = start_date
    all_end = end_date

    # ──────────────────────────────────────────────────────────
    # 🎯 [하드코딩 제거 1]: DB 설정값을 기준으로 휴가/연차류 자동 추출
    # deduction(차감)이 0보다 크거나, 근무일이 아닌데 이름이 '휴무'나 '정상'이 아닌 특수 근무만 자동 필터링!
    # ──────────────────────────────────────────────────────────
    leave_work_types = WorkType.objects.filter(
        Q(deduction__gt=0) | (Q(is_working_day=False) & ~Q(name__contains="휴무"))
    ).exclude(name__contains="정상")

    summary_schedules = DailySchedule.objects.filter(
        profile__in=profiles, date__range=(all_start, all_end),
        work_type__in=leave_work_types
    ).select_related('profile__process', 'work_type').order_by('date', 'profile__name')

    s_row = 2
    for s in summary_schedules:
        purpose_text = s.profile.user.reservation_set.filter(
            start_time__date=s.date
        ).first().title if s.profile.user.reservation_set.filter(start_time__date=s.date).exists() else "근무표 자동 연동"

        ws_summary.cell(row=s_row, column=1, value=s.profile.name)
        ws_summary.cell(row=s_row, column=2, value=s.profile.process.name if s.profile.process else "-")
        ws_summary.cell(row=s_row, column=3, value=s.date.strftime('%Y-%m-%d'))
        
        # ★ "연차"라고 치는 대신 DB의 이름을 그대로 가져옵니다. (예: 반차, 예비군공가 등 자동 반영)
        ws_summary.cell(row=s_row, column=4, value=s.work_type.name)
        ws_summary.cell(row=s_row, column=5, value=purpose_text)
        
        for c_idx in range(1, 6):
            c_cell = ws_summary.cell(row=s_row, column=c_idx)
            c_cell.font = Font(name='맑은 고딕', size=10)
            c_cell.border = thin_border
            c_cell.alignment = Alignment(horizontal='center', vertical='center')
        s_row += 1
        
    if s_row == 2:
        ws_summary.cell(row=2, column=1, value="해당 기간 내 사용된 연차/공가 내역이 없습니다.")
        ws_summary.merge_cells('A2:E2')
        ws_summary.cell(row=2, column=1).alignment = Alignment(horizontal='center')
    # ──────────────────────────────────────────────────────────

    # ──────────────────────────────────────────────────────────
    # 🎯 [하드코딩 제거 2]: DB의 근무 유형 색상을 그대로 읽어와 컬러맵 구성
    # ──────────────────────────────────────────────────────────
    wt_color_map = {}
    for wt in WorkType.objects.all():
        key_name = wt.short_name if wt.short_name else wt.name[:2]
        # openpyxl은 색상 코드에 '#'이 들어가면 에러가 나므로 없애줍니다. 색상이 비어있으면 하얀색 부여.
        bg_hex = wt.color.replace('#', '') if wt.color else 'FFFFFF'
        wt_color_map[key_name] = bg_hex

    # 시작월부터 종료월까지의 모든 (연도, 월) 쌍을 계산하여 리스트업
    months_to_progress = []
    cur_date = start_date.replace(day=1)
    while cur_date <= end_date:
        months_to_progress.append((cur_date.year, cur_date.month))
        if cur_date.month == 12:
            cur_date = cur_date.replace(year=cur_date.year + 1, month=1)
        else:
            cur_date = cur_date.replace(month=cur_date.month + 1)

    for y, m in months_to_progress:
        # 달력 일수 스캔 (루프 연도 y와 월 m 대입)
        _, num_days = calendar.monthrange(y, m)
        ws = wb.create_sheet(title=f"{y}년 {m}월 근무표")
        
        # 가로 타이틀 헤더 어펜드
        headers = ["이름 / 공정"]
        for day in range(1, num_days + 1):
            headers.append(f"{day}일")
        headers += ["출근", "휴무", "연차", "반차", "잔여 연차"]
        ws.append(headers)
        
        # 헤더 스타일링
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(bold=True, size=10)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = header_fill
            cell.border = thin_border
            
            # 주말 열 하이라이트 보정
            if 1 < col_num <= (num_days + 1):
                # ── [★ 수정]: target_year 대신 루프 연도 y 적용 ──
                d_check = date(y, m, col_num - 1)
                if d_check.weekday() == 5: cell.fill = sat_fill
                elif d_check.weekday() == 6: cell.fill = sun_fill

        # 해당 월의 모든 스케줄 셋 한방에 인메모리 로드 (쿼리 폭발 방지)
        # ── [★ 수정]: 상위 변수명 충돌 방지 및 루프 전용 연도 y 바인딩 ──
        start_month_date = date(y, m, 1)
        end_month_date = date(y, m, num_days)
        schedules = DailySchedule.objects.filter(
            profile__in=profiles, date__range=(start_month_date, end_month_date)
        ).select_related('work_type')

        # [profile_id][date_str] = work_type 맵 생성
        sched_map = {}
        for s in schedules:
            if s.profile_id not in sched_map: sched_map[s.profile_id] = {}
            sched_map[s.profile_id][s.date.strftime('%Y-%m-%d')] = s.work_type

        # 로우 데이터 채우기
        for p in profiles:
            row_cells = [f"{p.name}\n({p.process.name if p.process else '-'})"]
            
            # 카운팅 변수
            work_cnt, rest_cnt, leave_cnt, half_cnt = 0, 0, 0, 0
            
            for day in range(1, num_days + 1):
                # ── [★ 수정]: target_year 대신 y 적용 ──
                d_str = date(y, m, day).strftime('%Y-%m-%d')
                wt = sched_map.get(p.id, {}).get(d_str, None)
                
                if wt:
                    short_name = wt.short_name if wt.short_name else wt.name[:2]
                    row_cells.append(short_name)
                    
                    # 통계 누적
                    if wt.deduction == 1.0: leave_cnt += 1
                    elif 0 < wt.deduction < 1.0: half_cnt += 1
                    elif wt.is_working_day and wt.deduction == 0: work_cnt += 1
                    else:
                        if not wt.is_working_day: rest_cnt += 1
                else:
                    # 기본값 지정 규칙 (주말은 휴무, 평일은 F)
                    # ── [★ 수정]: target_year 대신 y 적용 ──
                    d_obj = date(y, m, day)
                    if d_obj.weekday() >= 5:
                        row_cells.append("휴무")
                        rest_cnt += 1
                    else:
                        row_cells.append("F")
                        work_cnt += 1
            
            # 통계 데이터 및 잔여 연차 마감 결합
            # ── [★ 수정]: target_year 대신 y 적용 ──
            total_leave = calculate_annual_leave_total(p, y)
            row_cells += [work_cnt, rest_cnt, leave_cnt, half_cnt, f"{total_leave - leave_cnt} / {total_leave}"]
            ws.append(row_cells)
            
            # 방금 추가된 로우 서식 지정
            curr_row = ws.max_row
            for col_num in range(1, len(row_cells) + 1):
                c = ws.cell(row=curr_row, column=col_num)
                c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                c.border = thin_border
                c.font = Font(size=10)
                
                # 가독성: 이름칸 왼쪽 정렬 보정
                if col_num == 1: c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
                
                # 특정 근무 유형 폰트 색상 하이라이트 (휴무는 빨간색 등)
                val = str(c.value)
                if "휴무" in val: 
                    c.font = Font(color="FF0000", size=10)
                
                # ──────────────────────────────────────────────────────────
                # 🎯 [하드코딩 제거 3]: 미리 만들어둔 wt_color_map에서 색상을 꺼내 자동으로 칠합니다!
                # ──────────────────────────────────────────────────────────
                elif val in wt_color_map:
                    dynamic_bg_hex = wt_color_map[val]
                    c.fill = PatternFill(start_color=dynamic_bg_hex, end_color=dynamic_bg_hex, fill_type='solid')
                    # 배경이 들어가면 글씨는 돋보이게 볼드(굵게) 처리
                    c.font = Font(bold=True, size=10, color="222222")

        # 셀 너비 자동 맞춤 최적화
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 7)
        ws.column_dimensions['A'].width = 15 # 이름 컬럼은 좀 더 넓게

    # 4. 파일 스트리밍 출력 반환
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    file_prefix = "year" if download_scope == "year" else f"{target_month}month"
    response['Content-Disposition'] = f'attachment; filename=pmtc_schedule_{target_year}_{file_prefix}.xlsx'
    wb.save(response)
    return response